from __future__ import annotations

import io

import pytest
from conftest import make_pdf

from clauselens.calculators import bond_cost, suggest_mode
from clauselens.gaps import EXPECTED, GAP_TEXT, detect_doc_type, find_gaps
from clauselens.ingest import MAX_BYTES, IngestError, extract_text
from clauselens.qa import ask, route_question
from clauselens.synth.generator import generate_documents

# ---------------------------------------------------------------- doc type / gaps


def test_doc_type_detection_on_synthetic_documents():
    docs = generate_documents(40, seed=4)
    hits = sum(detect_doc_type(d["text"])[0] == d["doc_type"] for d in docs)
    assert hits == len(docs)
    assert detect_doc_type("Hello world.") == ("unknown", 0.0)


def test_gap_detection_only_reports_predictable_expected_clauses():
    present = {"security_deposit", "lock_in_period"}
    predictable = set(EXPECTED["rent_agreement"]) - {"handover_inventory"}
    gaps = find_gaps("rent_agreement", present, predictable)
    assert "security_deposit" not in gaps and "rent_escalation" in gaps
    assert "handover_inventory" not in gaps  # classifier cannot predict it -> not reported
    assert find_gaps("unknown", set(), predictable) == []
    assert all(c in GAP_TEXT for ids in EXPECTED.values() for c in ids)


# ---------------------------------------------------------------- Q&A

CLAUSES = [
    {"index": 1, "kind": "clause", "number": "1.", "heading": "SECURITY DEPOSIT", "labels": ["security_deposit"],
     "text": "The Tenant paid Rs. 50,000/- as deposit, refundable within 30 days of vacating.",
     "explanation": {"facts_line": "Deposit: Rs. 50,000/- · Refund window: 30 days"}},
    {"index": 2, "kind": "clause", "number": "2.", "heading": "NOTICE PERIOD", "labels": ["notice_period"],
     "text": "Either party may end the tenancy by giving 60 days written notice.", "explanation": None},
    {"index": 0, "kind": "preamble", "text": "This agreement is made at Mumbai.", "labels": ["other"]},
]  # fmt: skip


def test_question_routing_english_and_hinglish():
    assert route_question("How much is the security deposit?") == ["security_deposit"]
    assert "notice_period" in route_question("Notice kitna dena hoga?")
    assert route_question("tell me a joke") == []


def test_answers_cite_the_matching_clause():
    result = ask(CLAUSES, "How much is the deposit and when is it refunded?", "en")
    assert result["found"] and result["citations"][0]["index"] == 1
    assert "Rs. 50,000/-" in result["answer"] and "Key terms: Deposit" in result["answer"]
    notice = ask(CLAUSES, "What is the notice period?", "hinglish")
    assert notice["citations"][0]["index"] == 2 and notice["answer"].startswith(
        "Agreement ke hisaab se"
    )


def test_unanswerable_questions_are_not_guessed():
    for question in ("What colour is the sky tonight?", "", "the of and"):
        result = ask(CLAUSES, question, "en")
        assert result["found"] is False and result["citations"] == []
    assert ask([], "deposit?", "en")["found"] is False


# ---------------------------------------------------------------- bond calculator


def test_bond_cost_modes():
    assert bond_cost(120000, 12, 3)["payable"] == 90000
    assert bond_cost(120000, 12, 3, "full")["payable"] == 120000
    assert bond_cost(120000, 12, 12, "full")["payable"] == 0  # bond completed
    assert bond_cost(120000, 12, 0)["remaining_months"] == 12
    with pytest.raises(ValueError):
        bond_cost(100, 0, 1)
    with pytest.raises(ValueError):
        bond_cost(100, 12, 1, "weird")
    assert suggest_mode(["full_amount_regardless"]) == "full" and suggest_mode([]) == "pro_rata"


# ---------------------------------------------------------------- ingest


def test_text_docx_and_pdf_extraction():
    from docx import Document

    assert extract_text("a.txt", b"1. Deposit Rs. 5,000/-") == "1. Deposit Rs. 5,000/-"
    doc = Document()
    doc.add_paragraph("1. SECURITY DEPOSIT: The Tenant pays a refundable deposit.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Rent", "Rs. 30,000"
    buf = io.BytesIO()
    doc.save(buf)
    text = extract_text("a.docx", buf.getvalue())
    assert "SECURITY DEPOSIT" in text and "Rent | Rs. 30,000" in text
    pdf = extract_text("a.pdf", make_pdf("The Tenant shall pay a security deposit of Rs. 40,000."))
    assert "security deposit" in pdf


def test_ingest_rejects_bad_files():
    with pytest.raises(IngestError, match="unsupported"):
        extract_text("a.exe", b"x")
    with pytest.raises(IngestError, match="no text"):
        extract_text("a.txt", b"  \n ")
    with pytest.raises(IngestError, match="larger"):
        extract_text("a.txt", b"x" * (MAX_BYTES + 1))
    with pytest.raises(IngestError, match="could not read"):
        extract_text("a.docx", b"this is not a docx")
    with pytest.raises(IngestError):
        extract_text("a.pdf", b"%PDF-1.4 broken")
