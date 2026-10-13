"""Turn an uploaded file (.txt / .md / .docx / .pdf) into plain text."""

from __future__ import annotations

import io
from pathlib import Path

MAX_BYTES = 5 * 1024 * 1024
MAX_CHARS = 200_000


class IngestError(ValueError):
    """Raised for files we cannot or will not read (message is safe to show to the user)."""


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _docx_text(data: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


_TESSERACT_LANGUAGES = {
    "af": "afr", "ar": "ara", "as": "asm", "bn": "ben", "ca": "cat",
    "cs": "ces", "da": "dan", "de": "deu", "el": "ell", "en": "eng",
    "eo": "epo", "es": "spa", "et": "est", "fa": "fas", "fi": "fin",
    "fr": "fra", "gu": "guj", "he": "heb", "hi": "hin", "hr": "hrv",
    "hu": "hun", "id": "ind", "it": "ita", "ja": "jpn", "kn": "kan",
    "ko": "kor", "lt": "lit", "lv": "lav", "ml": "mal", "mr": "mar",
    "ms": "msa", "ne": "nep", "nl": "nld", "no": "nor", "or": "ori",
    "pa": "pan", "pl": "pol", "pt": "por", "ro": "ron", "ru": "rus",
    "sk": "slk", "sl": "slv", "sq": "sqi", "sv": "swe", "ta": "tam",
    "te": "tel", "th": "tha", "tr": "tur", "uk": "ukr", "ur": "urd",
    "vi": "vie", "zh": "chi_sim",
}
_MAX_OCR_PAGES = 20


def _ocr_pdf_text(data: bytes, source_language: str = "auto") -> str:
    """OCR a scanned PDF when PyMuPDF, Pillow, Tesseract and trained data are installed."""
    try:
        import pymupdf
        import pytesseract
        from PIL import Image
    except ImportError as exc:
        raise IngestError(
            "This PDF appears to be scanned. Install the optional OCR packages with "
            '`python -m pip install -e ".[ocr]"`, install the Tesseract application, '
            "and add the required Tesseract language data. Then upload the PDF again."
        ) from exc

    try:
        available = set(pytesseract.get_languages(config=""))
        if source_language == "auto":
            # Mixed English + Indian-language agreements are common. Use a small set of
            # installed scripts for auto mode; callers can select a single source language for
            # better accuracy on a scanned contract.
            preferred = ["eng", "hin", "mar", "guj", "ben", "asm", "pan", "tam", "tel", "kan", "mal", "ori", "urd", "ara", "nep"]
            requested = [code for code in preferred if code in available][:4] or (["eng"] if "eng" in available else [])
        elif source_language == "hinglish":
            requested = [code for code in ("eng", "hin") if code in available]
        else:
            mapped_language = _TESSERACT_LANGUAGES.get(source_language)
            if mapped_language is None:
                raise IngestError(
                    f"Scanned-PDF OCR is not configured for language {source_language!r}. "
                    "Use a text-based file or configure its Tesseract language mapping."
                )
            # Keeping English enabled helps read bilingual legal documents where headings and
            # defined terms remain in English but the body is in a local language.
            requested = [code for code in (mapped_language, "eng") if code in available]
        languages = [code for code in requested if code in available]
        if not languages:
            needed = ", ".join(requested)
            raise IngestError(
                f"OCR language data {needed!r} is not installed in Tesseract. Install that "
                "language pack or choose a readable text-based PDF, DOCX, or TXT file."
            )
        language_arg = "+".join(languages)
        parts: list[str] = []
        with pymupdf.open(stream=data, filetype="pdf") as document:
            if len(document) > _MAX_OCR_PAGES:
                raise IngestError(
                    f"Scanned PDF OCR is limited to {_MAX_OCR_PAGES} pages per upload. "
                    "Split the document and upload each part separately."
                )
            for page in document:
                pixmap = page.get_pixmap(dpi=200, alpha=False)
                image = Image.open(io.BytesIO(pixmap.tobytes("png")))
                parts.append(pytesseract.image_to_string(image, lang=language_arg))
                if sum(len(part) for part in parts) >= MAX_CHARS:
                    break
        return "\n".join(parts).strip()
    except IngestError:
        raise
    except Exception as exc:
        # Never silently return a blank analysis for a scanned contract.
        raise IngestError(
            "OCR could not read this PDF. Check that the Tesseract application and the selected "
            "language data are installed, or provide a text-based PDF/DOCX/TXT file."
        ) from exc


def extract_text(filename: str, data: bytes, *, source_language: str = "auto") -> str:
    if len(data) > MAX_BYTES:
        raise IngestError(f"file is larger than {MAX_BYTES // (1024 * 1024)} MB")
    suffix = Path(filename or "").suffix.lower()
    try:
        if suffix in (".txt", ".md", ""):
            text = _decode(data)
        elif suffix == ".docx":
            text = _docx_text(data)
        elif suffix == ".pdf":
            text = _pdf_text(data)
        else:
            raise IngestError(f"unsupported file type {suffix!r}; use .txt, .docx or .pdf")
    except IngestError:
        raise
    except Exception as exc:  # corrupt files raise many different library errors
        raise IngestError(f"could not read {suffix or 'file'}: {exc.__class__.__name__}") from exc
    text = text.replace("\x00", "").strip()
    if suffix == ".pdf" and len(text) < 20:
        text = _ocr_pdf_text(data, source_language=source_language)
        if len(text) < 20:
            raise IngestError(
                "No readable text could be extracted from this PDF. Check its quality and "
                "the installed OCR language data."
            )
    if not text:
        raise IngestError("the file contains no text")
    return text[:MAX_CHARS]
